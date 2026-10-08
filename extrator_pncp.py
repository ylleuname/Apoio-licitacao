import cloudscraper
import pandas as pd
import os
import time
import requests
import re
from datetime import datetime
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


def baixar_arquivos_pncp(cnpj: str, ano: str, seq: str, orgao: str, scraper, pasta_raiz: str):
    orgao_limpo = re.sub(r'[\\/*?:"<>|]', "", str(orgao)).strip()[:60]
    if not orgao_limpo:
        orgao_limpo = f"Orgao_{cnpj}"

    nome_pasta = f"{orgao_limpo}_Seq_{seq}"
    pasta_destino = os.path.join(pasta_raiz, nome_pasta)

    url_lista_arquivos = f"https://pncp.gov.br/api/pncp/v1/orgaos/{cnpj}/compras/{ano}/{seq}/arquivos"

    try:
        res = scraper.get(url_lista_arquivos, headers={'Accept': 'application/json'}, timeout=15)
        if res.status_code != 200:
            return

        arquivos = res.json()
        if not arquivos:
            return

        os.makedirs(pasta_destino, exist_ok=True)

        for arq in arquivos:
            titulo = arq.get('tituloDocumento') or arq.get('titulo') or "Documento"
            url_down = arq.get('url') or arq.get('linkArquivo') or arq.get('urlDocumento')

            if not url_down:
                seq_arq = arq.get('sequencialDocumento') or arq.get('sequencialArquivo') or arq.get('sequencial')
                if seq_arq:
                    url_down = f"https://pncp.gov.br/api/pncp/v1/orgaos/{cnpj}/compras/{ano}/{seq}/arquivos/{seq_arq}"

            if not url_down:
                continue

            titulo_limpo = re.sub(r'[\\/*?:"<>|]', "", str(titulo)).strip()[:80]
            if not titulo_limpo:
                titulo_limpo = f"Documento_{seq_arq}"

            if not '.' in titulo_limpo[-5:]:
                titulo_limpo += ".pdf"

            caminho_arquivo = os.path.join(pasta_destino, titulo_limpo)
            if os.path.exists(caminho_arquivo):
                continue

            try:
                doc_res = scraper.get(url_down, stream=True, allow_redirects=True, timeout=60)
                if doc_res.status_code == 200:
                    content_type = doc_res.headers.get('Content-Type', '')
                    if 'application/json' in content_type:
                        continue
                    with open(caminho_arquivo, 'wb') as f:
                        for chunk in doc_res.iter_content(chunk_size=8192):
                            if chunk:
                                f.write(chunk)
            except Exception:
                pass
    except Exception:
        pass


def buscar_editais_pncp(palavras_chave: list, pasta_raiz: str, app_gui=None, baixar_arquivos=True) -> pd.DataFrame:
    base_url = "https://pncp.gov.br/api/search"
    resultados_totais = []
    hoje = datetime.now()

    scraper = cloudscraper.create_scraper(browser={'browser': 'chrome', 'platform': 'windows', 'desktop': True})
    headers_detalhe = {'Accept': 'application/json'}

    status_abertos = ['recebendo_proposta']

    for keyword in palavras_chave:
        for status_slug in status_abertos:
            pagina_atual = 1

            while True:
                if app_gui:
                    texto_status = f"A procurar: '{keyword}' (Página {pagina_atual})..." if pagina_atual > 1 else f"A procurar: '{keyword}'..."
                    app_gui.atualizar_status(texto_status)

                params = {
                    'q': keyword,
                    'tipos_documento': 'edital',
                    'status': status_slug,
                    'ordenacao': '-data',
                    'pagina': pagina_atual,
                    'tam_pagina': 50
                }

                max_tentativas = 3
                sucesso_requisicao = False
                dados = []

                for tentativa in range(1, max_tentativas + 1):
                    try:
                        response = scraper.get(base_url, params=params, timeout=60)
                        response.raise_for_status()
                        dados = response.json().get('items', [])
                        sucesso_requisicao = True
                        break
                    except:
                        if tentativa < max_tentativas:
                            time.sleep(5)

                if not sucesso_requisicao or not dados:
                    break

                for licitacao in dados:
                    id_licitacao = licitacao.get('numero_controle_pncp', 'N/A')
                    cnpj = licitacao.get('orgao_cnpj')
                    ano = licitacao.get('ano')
                    seq = licitacao.get('numero_sequencial')

                    data_fim_str = licitacao.get('data_fim_recebimento_propostas')

                    valor_total = licitacao.get('valor_total_estimado')

                    if (not data_fim_str or str(
                            data_fim_str).strip() == "" or valor_total is None) and cnpj and ano and seq:
                        detalhe_url = f"https://pncp.gov.br/api/pncp/v1/orgaos/{cnpj}/compras/{ano}/{seq}"
                        try:
                            res_detalhe = scraper.get(detalhe_url, headers=headers_detalhe, timeout=15)
                            if res_detalhe.status_code == 200:
                                json_detalhe = res_detalhe.json()

                                if not data_fim_str or str(data_fim_str).strip() == "":
                                    data_fim_str = (
                                            json_detalhe.get('dataHoraFimRecebimentoPropostas') or
                                            json_detalhe.get('dataFimRecebimentoPropostas') or
                                            json_detalhe.get('dataHoraAberturaProposta') or
                                            json_detalhe.get('dataAberturaProposta')
                                    )

                                if valor_total is None:
                                    valor_total = json_detalhe.get('valorTotalEstimado')
                        except:
                            pass

                    is_valid = True

                    if data_fim_str and str(data_fim_str).strip() != "":
                        try:
                            data_fim_dt = pd.to_datetime(data_fim_str).tz_localize(None)
                            if data_fim_dt < hoje:
                                is_valid = False
                        except:
                            pass

                    if not is_valid:
                        continue

                    orgao = licitacao.get('orgao_nome', 'Não Informado')

                    if valor_total is not None:
                        try:
                            valor_formatado = f"R$ {float(valor_total):,.2f}".replace(",", "X").replace(".",
                                                                                                        ",").replace(
                                "X", ".")
                        except ValueError:
                            valor_formatado = str(valor_total)
                    else:
                        valor_formatado = "Não Informado"

                    if baixar_arquivos and cnpj and ano and seq:
                        baixar_arquivos_pncp(str(cnpj), str(ano), str(seq), orgao, scraper, pasta_raiz)

                    objeto_texto = str(
                        licitacao.get('description') or licitacao.get('title') or licitacao.get(
                            'informacao_complementar') or ""
                    ).strip()

                    data_pub_str = licitacao.get('data_publicacao_pncp')
                    data_pub_formatada = pd.to_datetime(data_pub_str.split('T')[0]).strftime(
                        '%d/%m/%Y') if data_pub_str else 'N/A'

                    link_acesso = f"https://pncp.gov.br/app/editais/{cnpj}/{ano}/{seq}" if (
                            cnpj and ano and seq) else licitacao.get('itemUrl', 'Link indisponível')

                    resultados_totais.append({
                        'ID Licitação': id_licitacao,
                        'Órgão Nome': orgao,
                        'Objeto do Edital': objeto_texto,
                        'Data Publicação': data_pub_formatada,
                        'Valor Total Estimado': valor_formatado,
                        'Link Acesso': link_acesso
                    })

                pagina_atual += 1

    df = pd.DataFrame(resultados_totais)
    if not df.empty:
        df = df.drop_duplicates(subset=['ID Licitação'])

    return df


def salvar_em_excel(df_editais: pd.DataFrame, pasta_raiz: str):
    nome_arquivo = os.path.join(pasta_raiz, "editais_monitorados.xlsx")
    if df_editais.empty:
        return 0

    # Criação estruturada do Excel
    with pd.ExcelWriter(nome_arquivo, engine='openpyxl') as writer:
        df_editais.to_excel(writer, index=False, sheet_name='Oportunidades')

        workbook = writer.book
        worksheet = writer.sheets['Oportunidades']

        # Definição do Design (Cor Bloko #3751A3 no formato Hex do OpenPyXL - sem a hashtag)
        header_fill = PatternFill(start_color="3751A3", end_color="3751A3", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        wrap_alignment = Alignment(wrap_text=True, vertical="top")
        center_alignment = Alignment(horizontal="center", vertical="center")

        # Aplica os estilos no cabeçalho e define a largura das colunas
        for col_idx, col in enumerate(df_editais.columns, 1):
            col_letter = get_column_letter(col_idx)

            # Pinta e formata o título
            header_cell = worksheet[f"{col_letter}1"]
            header_cell.fill = header_fill
            header_cell.font = header_font
            header_cell.alignment = center_alignment

            # Trava a largura da coluna (~277 pixels)
            worksheet.column_dimensions[col_letter].width = 38

            # Ativa a quebra de texto (wrap text) e o alinhamento para o topo no restante das linhas
            for row_idx in range(2, len(df_editais) + 2):
                worksheet[f"{col_letter}{row_idx}"].alignment = wrap_alignment

    return len(df_editais)